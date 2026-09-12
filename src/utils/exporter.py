import openpyxl
from docx import Document
from docx.shared import Inches
import io

def export_results(choice, extracted_results, show_lines, show_acc, path):
    ext_map = {
        "Export as .txt": ".txt",
        "Export as .md": ".md",
        "Export as .docx": ".docx",
        "Export as .xlsx": ".xlsx"
    }
    ext = ext_map.get(choice, ".txt")

    try:
        if ext == ".txt":
            with open(path, "w", encoding="utf-8") as f:
                for item in extracted_results:
                    if item['type'] == 'text':
                        line = (f"{item['line']}. " if show_lines else "") + item['text']
                        if show_acc: line += f"  [Acc: {item['accuracy']:.1f}%]"
                        f.write(line + "\n")

        elif ext == ".md":
            with open(path, "w", encoding="utf-8") as f:
                f.write("# Extracted Khmer Text\n\n")
                for item in extracted_results:
                    if item['type'] == 'text':
                        line = (f"{item['line']}. " if show_lines else "") + item['text']
                        if show_acc: line += f"  [Acc: {item['accuracy']:.1f}%]"
                        f.write(line + "\n")
                    elif item['type'] == 'image':
                        f.write("\n![Logo](logo_in_memory)\n")

        elif ext == ".docx":
            doc = Document()
            doc.add_heading('Extracted Khmer Text', 0)
            for item in extracted_results:
                if item['type'] == 'text':
                    line = (f"{item['line']}. " if show_lines else "") + item['text']
                    if show_acc: line += f"  [Acc: {item['accuracy']:.1f}%]"
                    doc.add_paragraph(line)
                elif item['type'] == 'image' and 'image_obj' in item:
                    # PERFORMANCE FIX: Embed in-memory PIL image directly into DOCX via BytesIO
                    img_byte_arr = io.BytesIO()
                    item['image_obj'].save(img_byte_arr, format='PNG')
                    img_byte_arr.seek(0)
                    doc.add_paragraph().add_run().add_picture(img_byte_arr, width=Inches(3.0))
            doc.save(path)

        elif ext == ".xlsx":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "OCR Results"
            headers = (["Line #"] if show_lines else []) + ["Text"] + (["Accuracy (%)"] if show_acc else [])
            ws.append(headers)
            for item in extracted_results:
                if item['type'] == 'text':
                    row = ([item['line']] if show_lines else []) + [item['text']] + (
                        [round(item['accuracy'], 1)] if show_acc else [])
                    ws.append(row)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(
                    max(len(str(cell.value)) for cell in col if cell.value) + 4, 100)
            wb.save(path)

        return True, path
    except Exception as e:
        return False, str(e)